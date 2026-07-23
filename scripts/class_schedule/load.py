"""
Load De Anza class-schedule section rows into a single DynamoDB table.

Both the 2026-27 published schedule (Excel) and the 2025-26 historic schedule
with enrollment (CSV) live in ONE table, distinguished by term_code. This lets
the conflict analysis pull an entire quarter with one query and compare a
planned section against last year's enrollment.

Table key:
    partition  term_code    e.g. "202722"       (a whole quarter = one partition)
    sort       section_key  "<CRN>#<seq>"        (seq disambiguates multi-meeting rows)

Usage:
    # Create the table (one time) and load all four 2026-27 Excel files
    python scripts/class_schedule/load.py --create-table

    # Load the 2025-26 enrollment CSVs from a local dir, tagging the year
    python scripts/class_schedule/load.py --format csv \
        --dir data/2025-26-class-schedule --academic-year 2025-26

    # Read the Excel files straight from S3 instead of locally
    python scripts/class_schedule/load.py --from-s3 --s3-prefix data/2026-27-class-schedule/
"""

import io
import csv
import sys
import glob
import argparse
from decimal import Decimal
from pathlib import Path

import openpyxl
from botocore.exceptions import ClientError

# Make the shared helpers importable when run directly
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from common import (
    PROJECT_ROOT, DEFAULT_BUCKET,
    get_session, get_s3_client, get_dynamodb, create_table_if_absent,
)

DEFAULT_TABLE = "deanza-class-schedule"
DEFAULT_LOCAL_DIR = PROJECT_ROOT / "data" / "2026-27-class-schedule"

# Columns worth promoting to named attributes (source header -> item field).
# Anything not listed still rides along in a generic "raw" map so nothing is lost.
COLUMN_MAP = {
    "Term_Code": "term_code",
    "Division": "division",
    "Subject": "subject",
    "Number": "number",
    "Section": "section",
    "CRN": "crn",
    "InstructMethCode": "instruction_method",
    "SectionStat": "section_status",
    "PartTermCode": "part_term",
    "Start_Date": "start_date",
    "End_Date": "end_date",
    "MeetingType": "meeting_type",
    "Meeting_Days": "meeting_days",
    "Meeting_Times": "meeting_times",
    "Room": "room",
    "Building": "building",
    "MaxEnroll": "max_enroll",
    "TotalEnroll": "total_enroll",
    "WaitlistCapacity": "waitlist_capacity",
    "WaitlistCount": "waitlist_count",
    "Units": "units",
    "LastName": "instructor_last",
    "FirstName": "instructor_first",
}


# ---- source readers: each yields plain dicts keyed by the source headers ----

def read_xlsx_bytes(data):
    """Yield row dicts from an xlsx byte stream (first row = headers)."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = None
    for row in ws.iter_rows(values_only=True):
        if header is None:
            header = list(row)
            continue
        yield dict(zip(header, row))
    wb.close()


def read_csv_bytes(data):
    """Yield row dicts from a CSV byte stream.

    Assumes the 2025-26 CSVs use the same SIS column headers as the Excel export
    (Term_Code, Subject, CRN, ...). Verify against the first pulled file; if the
    headers differ, extend COLUMN_MAP rather than changing this reader.
    """
    text = io.StringIO(data.decode("utf-8-sig"))
    yield from csv.DictReader(text)


def iter_local_files(directory, fmt):
    """Yield (name, bytes) for every matching file in a local directory."""
    ext = "xlsx" if fmt == "xlsx" else "csv"
    for path in sorted(glob.glob(str(Path(directory) / f"*.{ext}"))):
        with open(path, "rb") as f:
            yield Path(path).name, f.read()


def iter_s3_files(session, prefix, fmt):
    """Yield (key, bytes) for every matching object under an S3 prefix."""
    ext = ".xlsx" if fmt == "xlsx" else ".csv"
    s3 = get_s3_client(session)
    paginator = s3.get_paginator("list_objects_v2")
    for page in paginator.paginate(Bucket=DEFAULT_BUCKET, Prefix=prefix):
        for obj in page.get("Contents", []):
            key = obj["Key"]
            if key.lower().endswith(ext):
                body = s3.get_object(Bucket=DEFAULT_BUCKET, Key=key)["Body"].read()
                yield key, body


# ---- value coercion for DynamoDB ----

def num(value):
    """Coerce a numeric-ish value to int/Decimal (DynamoDB rejects float)."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    try:
        # Decimal via str keeps 4.5 exact and satisfies boto3's number type
        return Decimal(str(value))
    except Exception:
        return None


def as_str(value):
    """Stringify dates/numbers/None into clean strings for text attributes."""
    if value is None:
        return ""
    # datetime and date both expose isoformat()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value).strip()


def to_item(row, academic_year, seq, source_file=""):
    """Shape one source row into a DynamoDB item.

    term_code + "<CRN>#<seq>" form the composite key; numeric enrollment fields
    are coerced to numbers so they can be summed/compared, and everything else
    is stored as trimmed strings. Unmapped source columns are kept under "raw".
    source_file records the originating file, since one load can span several
    quarters/years and the academic_year label is only a coarse dataset tag.
    """
    term_code = as_str(row.get("Term_Code"))
    crn = as_str(row.get("CRN")) or "NOCRN"

    subject = as_str(row.get("Subject"))
    number = as_str(row.get("Number"))

    numeric_fields = {"max_enroll", "total_enroll", "waitlist_capacity",
                      "waitlist_count", "units"}

    item = {
        "term_code": term_code,
        "section_key": f"{crn}#{seq}",
        "academic_year": academic_year,
        "source_file": source_file,
        # A normalized course label helps match against pathway course codes
        "course": f"{subject} {number}".strip(),
        "raw": {},
    }

    # Promote mapped columns; coerce the numeric ones, stringify the rest
    for src, field in COLUMN_MAP.items():
        val = row.get(src)
        item[field] = num(val) if field in numeric_fields else as_str(val)

    # Convenience: combined instructor name
    item["instructor"] = f"{item.get('instructor_first','')} {item.get('instructor_last','')}".strip()

    # Preserve any columns we didn't explicitly map (as strings)
    for src, val in row.items():
        if src not in COLUMN_MAP and src not in ("Term_Code", "CRN"):
            item["raw"][as_str(src)] = as_str(val)

    return item


def load(table, files, reader, academic_year):
    """Write every row from every file, assigning a per-CRN sequence number.

    The sequence disambiguates sections that occupy multiple meeting rows
    (e.g. a lecture + lab under one CRN) so they don't overwrite each other.
    """
    total = 0
    with table.batch_writer() as batch:
        for name, data in files:
            seen = {}  # (term_code, crn) -> next seq
            rows = 0
            for row in reader(data):
                if not as_str(row.get("Term_Code")):
                    continue  # skip blank/footer rows
                crn = as_str(row.get("CRN")) or "NOCRN"
                key = (as_str(row.get("Term_Code")), crn)
                seq = seen.get(key, 0)
                seen[key] = seq + 1
                batch.put_item(Item=to_item(row, academic_year, seq, source_file=name))
                rows += 1
            total += rows
            print(f"  {name}: {rows} rows")
    return total


def main():
    parser = argparse.ArgumentParser(description="Load class schedule into DynamoDB")
    parser.add_argument("--table", default=DEFAULT_TABLE, help="DynamoDB table name")
    parser.add_argument("--create-table", action="store_true",
                        help="Create the table first if it does not exist")
    parser.add_argument("--format", choices=["xlsx", "csv"], default="xlsx",
                        help="Source file format (default: xlsx)")
    parser.add_argument("--dir", default=str(DEFAULT_LOCAL_DIR),
                        help="Local directory of schedule files")
    parser.add_argument("--from-s3", action="store_true",
                        help="Read files from S3 instead of the local directory")
    parser.add_argument("--s3-prefix", default="data/2026-27-class-schedule/",
                        help="S3 key prefix to read schedule files from")
    parser.add_argument("--academic-year", default="2026-27",
                        help="Academic-year label stored on each item")
    args = parser.parse_args()

    session = get_session()
    dynamodb = get_dynamodb(session)

    if args.create_table:
        created = create_table_if_absent(args.table, session)
        print(f"{'Created' if created else 'Already exists'}: {args.table}")

    # Pick the source (S3 vs local) and the matching byte reader
    if args.from_s3:
        files = iter_s3_files(session, args.s3_prefix, args.format)
        print(f"Reading {args.format} files from s3://{DEFAULT_BUCKET}/{args.s3_prefix}")
    else:
        files = iter_local_files(args.dir, args.format)
        print(f"Reading {args.format} files from {args.dir}")

    reader = read_xlsx_bytes if args.format == "xlsx" else read_csv_bytes

    table = dynamodb.Table(args.table)
    try:
        total = load(table, files, reader, args.academic_year)
    except ClientError as e:
        if e.response["Error"]["Code"] == "ResourceNotFoundException":
            print(f"Table '{args.table}' not found. Re-run with --create-table first.")
            return
        raise

    print(f"Done — wrote {total} section rows to '{args.table}'.")


if __name__ == "__main__":
    main()
