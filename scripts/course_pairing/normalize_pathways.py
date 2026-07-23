"""
Transform the pathway maps into cleaned, canonical course lists.

Reads the source pathways (local JSON by default, or the `deanza-pathways`
table with --from-table), normalizes every quarter's course text via the pure
pathway_parser, and reports coverage. Writing the result to the separate
`deanza-pathways-normalized` table is gated behind --apply (see Task 3); by
default this is a dry run that mutates nothing.

The original `deanza-pathways` table and source JSON are never modified.

Usage:
    # Dry run against local JSON (no AWS writes; prints coverage + samples)
    python scripts/course_pairing/normalize_pathways.py

    # Dry run reading pathways from the source table instead of local JSON
    python scripts/course_pairing/normalize_pathways.py --from-table
"""

import io
import sys
import glob
import json
import argparse
from pathlib import Path

# Make shared helpers and the course_pairing package importable when run directly.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from common import PROJECT_ROOT, get_table, get_dynamodb, create_table_if_absent  # noqa: E402
from course_pairing.pathway_parser import normalize_pathway_years  # noqa: E402

SOURCE_TABLE = "deanza-pathways"
NORMALIZED_TABLE = "deanza-pathways-normalized"
SCHEDULE_TABLE = "deanza-class-schedule"
NORMALIZED_VERSION = 1

DEFAULT_JSON = (
    PROJECT_ROOT / "data" / "s3-raw" / "data" / "program-pathways" / "deanza_pathways.json"
)
LOCAL_SCHEDULE_DIR = PROJECT_ROOT / "data" / "2026-27-class-schedule"

# Metadata carried onto the normalized items (identity, not source course text).
META_FIELDS = ("program_name", "village", "credential_type", "source_file", "page_number")


# ---- subject allow-list (what anchors a real course during parsing) ----

def subjects_from_table(table_name=SCHEDULE_TABLE):
    """Distinct uppercase subject codes from the schedule table (projection scan)."""
    table = get_table(table_name)
    subjects = set()
    kwargs = {"ProjectionExpression": "subject"}
    while True:
        resp = table.scan(**kwargs)
        for item in resp.get("Items", []):
            s = str(item.get("subject", "")).strip().upper()
            if s:
                subjects.add(s)
        if "LastEvaluatedKey" not in resp:
            break
        kwargs["ExclusiveStartKey"] = resp["LastEvaluatedKey"]
    return subjects


def subjects_from_local_schedule(directory=LOCAL_SCHEDULE_DIR):
    """Fallback subject list read from the local 2026-27 xlsx schedule files.

    Used when the schedule table isn't reachable (e.g. expired STS creds) so a
    dry run still works offline. Covers only the local files, so the final load
    (Task 6) should prefer the table for full-catalog subject coverage.
    """
    import openpyxl

    subjects = set()
    for path in sorted(glob.glob(str(Path(directory) / "*.xlsx"))):
        with open(path, "rb") as f:
            wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = list(next(rows))
        idx = header.index("Subject") if "Subject" in header else None
        if idx is not None:
            for row in rows:
                if idx < len(row) and row[idx]:
                    subjects.add(str(row[idx]).strip().upper())
        wb.close()
    return subjects


def load_subjects(prefer_table=True):
    """Return the subject allow-list, preferring the table, falling back local."""
    if prefer_table:
        try:
            subjects = subjects_from_table()
            if subjects:
                print(f"Subject allow-list: {len(subjects)} subjects from schedule table.")
                return subjects
        except Exception as e:  # noqa: BLE001 - want a clean offline fallback
            print(f"Could not read schedule table ({type(e).__name__}); "
                  f"falling back to local schedule files.")
    subjects = subjects_from_local_schedule()
    print(f"Subject allow-list: {len(subjects)} subjects from local schedule files.")
    return subjects


# ---- source pathways ----

def load_source_pathways(from_table, json_path):
    """Load raw pathway records from the source table or the local JSON."""
    if from_table:
        table = get_table(SOURCE_TABLE)
        items = []
        kwargs = {}
        while True:
            resp = table.scan(**kwargs)
            items.extend(resp.get("Items", []))
            if "LastEvaluatedKey" not in resp:
                break
            kwargs["ExclusiveStartKey"] = resp["LastEvaluatedKey"]
        return items
    with open(json_path) as f:
        return json.load(f)["pathways"]


def pathway_id(record):
    """The normalized table's key: same identity rule as the source table."""
    if record.get("pathway_id"):
        return record["pathway_id"]
    return f"{record.get('source_file', '')}#{record.get('page_number', '')}"


def build_normalized_item(record, subjects):
    """Shape one source pathway into a normalized-table item (no source text)."""
    item = {"pathway_id": pathway_id(record), "normalization_version": NORMALIZED_VERSION}
    for field in META_FIELDS:
        if record.get(field) is not None:
            item[field] = record[field]
    item["years"] = normalize_pathway_years(record.get("years", {}), subjects)
    return item


def build_all(records, subjects):
    return [build_normalized_item(r, subjects) for r in records]


# ---- coverage reporting ----

def summarize(items):
    """Aggregate coverage stats across all normalized items for the dry run."""
    quarters = matched_quarters = 0
    total_courses = total_unresolved = 0
    distinct_courses = set()
    for item in items:
        for quarters_map in item["years"].values():
            for q in quarters_map.values():
                quarters += 1
                nc = q["normalized_courses"]
                total_courses += len(nc)
                total_unresolved += len(q["unresolved_entries"])
                distinct_courses.update(nc)
                if nc:
                    matched_quarters += 1
    return {
        "pathways": len(items),
        "quarters": quarters,
        "quarters_with_courses": matched_quarters,
        "total_course_slots": total_courses,
        "distinct_courses": len(distinct_courses),
        "unresolved_entries": total_unresolved,
    }


# ---- writing to the normalized table (never the source) ----

def write_items(table, items):
    """Batch-write items to the destination table (idempotent by pathway_id)."""
    count = 0
    with table.batch_writer() as batch:
        for item in items:
            batch.put_item(Item=item)
            count += 1
    return count


def apply_to_table(items, table_name, create_table):
    """Guarded write path. Refuses to touch the source pathways table."""
    if table_name == SOURCE_TABLE:
        raise ValueError(
            f"Refusing to write to the source table '{SOURCE_TABLE}'. "
            f"Normalized data goes to '{NORMALIZED_TABLE}'."
        )
    if create_table:
        created = create_table_if_absent(table_name)
        print(f"{'Created' if created else 'Already exists'}: {table_name}")
    table = get_dynamodb().Table(table_name)
    written = write_items(table, items)
    print(f"Wrote {written} items to '{table_name}' (re-runnable: overwrites by key).")
    return written


def print_report(items, stats, samples):
    header = "=== Normalization coverage ==="
    print(f"\n{header}")
    for k, v in stats.items():
        print(f"  {k:24} {v}")

    print(f"\n=== {samples} sample normalized pathways ===")
    for item in items[:samples]:
        print(f"\n- {item['pathway_id']}  |  {item.get('program_name','')}")
        for yk, quarters in item["years"].items():
            for qk, q in quarters.items():
                if q["normalized_courses"] or q["unresolved_entries"]:
                    print(f"    {yk}/{qk}: courses={q['normalized_courses']} "
                          f"unresolved={q['unresolved_entries']}")


def main():
    parser = argparse.ArgumentParser(
        description="Normalize pathway course lists into a new table (dry run by default)")
    parser.add_argument("--from-table", action="store_true",
                        help="Read source pathways from deanza-pathways instead of local JSON")
    parser.add_argument("--json", default=str(DEFAULT_JSON), help="Local pathways JSON path")
    parser.add_argument("--table", default=NORMALIZED_TABLE,
                        help="Destination table (never the source table)")
    parser.add_argument("--apply", action="store_true",
                        help="Write to the destination table. Without this it's a dry run.")
    parser.add_argument("--create-table", action="store_true",
                        help="Create the destination table first if it does not exist")
    parser.add_argument("--samples", type=int, default=3, help="How many sample pathways to print")
    args = parser.parse_args()

    subjects = load_subjects()
    records = load_source_pathways(args.from_table, args.json)
    print(f"Loaded {len(records)} source pathways "
          f"({'table' if args.from_table else 'local JSON'}).")

    items = build_all(records, subjects)
    print_report(items, summarize(items), args.samples)

    if args.apply:
        apply_to_table(items, args.table, args.create_table)
    else:
        print(f"\nDry run — nothing written. Re-run with --apply "
              f"(and --create-table the first time) to load '{args.table}'.")


if __name__ == "__main__":
    main()
